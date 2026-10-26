"""Build the IT&I(s)-2026 regular-presentation .docx from content.py.

Usage:
    .venv/bin/python papers/iti_2026/build_docx.py

Starts from the official conference template (journal_rules/ITIS-Word-
template.docx), strips its body paragraphs and tables, and appends the
approved text from content.py. This keeps the template's page setup, its
mirrored running header (title + page number) and its `Normal` style
(Times New Roman 14 pt) exactly as supplied.

Verification:
    /Applications/LibreOffice.app/Contents/MacOS/soffice --headless \\
        --convert-to pdf --outdir papers/iti_2026 papers/iti_2026/iti_2026_lapin.docx
    /opt/homebrew/bin/pdfinfo papers/iti_2026/iti_2026_lapin.pdf | grep Pages
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import content as C

HERE = Path(__file__).parent
TEMPLATE = HERE / "journal_rules" / "ITIS-Word-template.docx"
OUT = HERE / "iti_2026_lapin.docx"

FONT = "Times New Roman"
INDENT = Cm(1.25)


def _el(tag: str, **attrs) -> OxmlElement:
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), value)
    return node


def _replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    """Splice `new` in for `old` across a paragraph's runs, wherever the
    boundary falls. The template splits "IT&Is-2025" as separate runs
    ("...202", "5", ")"), so a single-run substring match would miss it."""
    runs = paragraph.runs
    full_text = "".join(run.text for run in runs)
    start = full_text.find(old)
    if start == -1:
        return False
    end = start + len(old)
    inserted = False
    pos = 0
    for run in runs:
        run_start, run_end = pos, pos + len(run.text)
        pos = run_end
        if run_end <= start or run_start >= end:
            continue  # outside the matched range, leave untouched
        local_start = max(start, run_start) - run_start
        local_end = min(end, run_end) - run_start
        before, after = run.text[:local_start], run.text[local_end:]
        run.text = before + (new if not inserted else "") + after
        inserted = True
    return True


def fix_header_year(doc: Document, section) -> None:
    """The template header text is a literal string, not a field; the
    conference year in it is stale. Rewrite it in both mirrored headers,
    leaving the tab spacing and the PAGE field runs untouched."""
    old, new = "IT&Is-2025", "IT&I(s)-2026"
    for header in (section.header, section.even_page_header):
        for paragraph in header.paragraphs:
            _replace_in_paragraph(paragraph, old, new)


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = _el("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def clear_body(doc: Document) -> None:
    """Remove every template paragraph and table; keep the trailing sectPr
    so the page setup and header references stay untouched."""
    for table in doc.tables:
        table._element.getparent().remove(table._element)
    for paragraph in doc.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def _para(doc: Document, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.alignment = align
    if indent is not None:
        fmt.first_line_indent = indent
    return p


def _run(paragraph, text: str, *, size=14, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _el("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    return run


def author_block(doc: Document) -> None:
    author_affiliations = getattr(C, "AUTHOR_AFFILIATION", {})
    for name, position in C.AUTHORS:
        p = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
        _run(p, name, bold=True, italic=True)
        p2 = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
        _run(p2, position)
        if name in author_affiliations:
            p3 = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
            _run(p3, author_affiliations[name], italic=True)
    affiliation = getattr(C, "AFFILIATION", "")
    if affiliation:
        p4 = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
        _run(p4, affiliation, italic=True)


def title_block(doc: Document) -> None:
    _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)  # empty line
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, C.TITLE, bold=True)


def body_block(doc: Document) -> None:
    _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)  # empty line before body
    for head, text in C.BODY:
        p = _para(doc, indent=INDENT)
        _run(p, head, bold=True)
        _run(p, " " + text)


def references_block(doc: Document) -> None:
    _para(doc)
    label = _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
    _run(label, "References:", size=12, bold=True)
    for index, reference in enumerate(C.REFERENCES, start=1):
        rp = _para(doc, indent=INDENT)
        _run(rp, f"{index}. {reference}", size=12)


def main() -> None:
    doc = Document(TEMPLATE)
    section = doc.sections[0]
    fix_header_year(doc, section)
    set_page_number_start(section, 1)
    clear_body(doc)

    author_block(doc)
    title_block(doc)
    body_block(doc)
    references_block(doc)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
